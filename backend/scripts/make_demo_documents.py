"""Generate sample files for the file types added on 2026-09-20 — a CSV,
an Excel workbook, a PowerPoint deck, an HTML page, a Python script and
an RTF memo — so the new extractors can be tried in the live app.

Default target: ~/Desktop/IntelliFile-Search-Demo (same as the other
download_demo_*.py scripts). Pass a different folder as the first
argument. Existing files with the same names are overwritten.

    backend/venv/bin/python backend/scripts/make_demo_documents.py [target]
"""

import sys
from pathlib import Path

DEFAULT_TARGET = Path.home() / "Desktop" / "IntelliFile-Search-Demo"


def main() -> None:
    target = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else DEFAULT_TARGET
    target.mkdir(parents=True, exist_ok=True)

    (target / "monthly expenses.csv").write_text(
        "month,category,amount,note\n"
        "January,rent,1200,paid on the 3rd\n"
        "January,groceries,340,mostly farmers market\n"
        "February,gym membership,45,annual plan renewed\n"
        "February,electricity,88,winter heating\n"
        "March,car insurance,210,six month premium\n"
    )

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Trip budget"
    ws.append(["item", "estimated", "actual"])
    ws.append(["flights to Lisbon", 420, 398])
    ws.append(["hostel four nights", 260, 260])
    ws.append(["surf lesson", 60, 75])
    ws.append(["total", "=SUM(B2:B4)", "=SUM(C2:C4)"])
    packing = wb.create_sheet("Packing list")
    for row in (["sunscreen"], ["passport"], ["rash guard"], ["camera charger"]):
        packing.append(row)
    wb.save(str(target / "lisbon trip.xlsx"))

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Community Garden Proposal"
    s1.placeholders[1].text = "Turning the empty lot on Elm Street into shared allotments"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Why now"
    s2.placeholders[1].text = "Forty families on the waiting list\nCity grant closes in October\nSoil test came back clean"
    s2.notes_slide.notes_text_frame.text = "Mention that the hardware store offered to donate tools."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Budget"
    box = s3.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    box.text = "Raised beds 1800, water line 950, fencing 600, compost bins 300"
    prs.save(str(target / "community garden pitch.pptx"))

    (target / "sourdough starter guide.html").write_text(
        "<html><head><title>Sourdough starter guide</title>"
        "<style>body{font-family:serif}</style></head><body>"
        "<script>console.log('tracking')</script>"
        "<h1>Keeping a sourdough starter alive</h1>"
        "<p>Feed it equal weights of flour and water every twelve hours at room temperature.</p>"
        "<p>If it smells like nail polish it is hungry, not dead.</p>"
        "</body></html>"
    )

    (target / "backup_photos.py").write_text(
        '"""Copy new photos from the camera card to the NAS, skipping duplicates by hash."""\n'
        "import hashlib\n"
        "import shutil\n"
        "from pathlib import Path\n\n"
        "CARD = Path('/Volumes/EOS_DIGITAL/DCIM')\n"
        "NAS = Path('/Volumes/home/photos')\n\n\n"
        "def file_hash(path: Path) -> str:\n"
        "    return hashlib.sha256(path.read_bytes()).hexdigest()\n\n\n"
        "def sync():\n"
        "    # skip anything already on the NAS so a re-run is safe\n"
        "    seen = {file_hash(p) for p in NAS.rglob('*.jpg')}\n"
        "    for photo in CARD.rglob('*.jpg'):\n"
        "        if file_hash(photo) not in seen:\n"
        "            shutil.copy2(photo, NAS / photo.name)\n"
    )

    (target / "landlord letter.rtf").write_text(
        r"{\rtf1\ansi{\fonttbl{\f0 Helvetica;}}\f0\fs24 "
        r"Dear Mr Okafor,\par "
        r"The kitchen tap has been dripping since the second week of March and the bathroom "
        r"extractor fan no longer switches on. Please arrange a repair visit.\par "
        r"Kind regards, Abhishek}"
    )

    print(f"Wrote 6 sample documents to {target}")
    print("Try: 'gym membership', 'surf lesson', 'garden grant october', 'starter smells', 'skip duplicates by hash', 'dripping tap'")


if __name__ == "__main__":
    main()
