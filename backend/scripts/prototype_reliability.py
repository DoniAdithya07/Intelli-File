"""Phase 11 regression: reliability and the file lifecycle.

- corrupt / unsupported / odd files (bad PDF, bad DOCX, bad XLSX, bad
  PPTX, zero bytes, binary disguised as .txt, a 3 MB single line) never
  abort a scan: they are counted, listed with a reason, and the good
  files around them are indexed;
- a file that vanishes between the walk and the read is skipped;
- an unreadable file (permissions) is a counted failure, and becomes
  indexed on the next scan once readable;
- a crash mid-index (record written, chunks not) is repaired by the next
  scan without rescanning anything else.

Run with:  backend/venv/bin/python backend/scripts/prototype_reliability.py
"""

import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.embeddings.model import EmbeddingModel, default_model_dir  # noqa: E402
from app.indexing import Indexer, index_folder  # noqa: E402
from app.indexing import folder_scan  # noqa: E402
from app.search import SearchService  # noqa: E402
from app.storage import FileRecordStore, KeywordStore, LanceDBVectorStore  # noqa: E402

MODEL_DIR = default_model_dir(Path(__file__).resolve().parents[1] / "models")


def set_readable(path: Path, readable: bool) -> None:
    """Make a file unreadable to the current user, and back. POSIX: mode 0.
    Windows: chmod only toggles the read-only flag, so an ACL deny-read
    entry for the current user does it (a real Windows ACL refusal, which
    is what the indexer must respect)."""
    if sys.platform == "win32":
        user = os.environ.get("USERNAME", "")
        args = ["/deny", f"{user}:(R)"] if not readable else ["/remove:d", user]
        subprocess.run(["icacls", str(path), *args], check=True, capture_output=True)
    else:
        os.chmod(path, 0o644 if readable else 0)


def running_as_root() -> bool:
    return hasattr(os, "geteuid") and os.geteuid() == 0


def main() -> None:
    workdir = Path(tempfile.mkdtemp())
    try:
        folder = workdir / "docs"
        folder.mkdir()
        good = {
            "good_one.txt": "The quarterly budget meeting moved to Thursday because the projector broke.",
            "good_two.md": "Sourdough needs a long cold ferment for flavour; bake at 230 C in a dutch oven.",
            "good_three.txt": "Car insurance renewal is due on 30 September; the excess is 350 euros.",
        }
        for name, text in good.items():
            (folder / name).write_text(text)
        bad = {
            "broken.pdf": b"%PDF-1.4 garbage \x00\x01\x02 not really a pdf",
            "broken.docx": b"PK\x03\x04 this is not a zip",
            "broken.xlsx": b"PK\x03\x04 nor is this",
            "broken.pptx": b"\x00\x00\x00 slides? no",
            "empty.txt": b"",
            "binary_disguised.txt": bytes(range(256)) * 200,
            "csv_injection.csv": b"name,amount\n=cmd|' /C calc'!A0,10\n=HYPERLINK(\"http://evil\"),20\n",
        }
        for name, data in bad.items():
            (folder / name).write_bytes(data)
        (folder / "one_long_line.txt").write_text("x" * (3 * 1024 * 1024))

        model = EmbeddingModel(MODEL_DIR)
        vector_store = LanceDBVectorStore(str(workdir / "vectors"))
        keyword_store = KeywordStore(workdir / "keyword.db")
        record_store = FileRecordStore(workdir / "files.db")
        indexer = Indexer(model, vector_store, keyword_store, record_store)
        search = SearchService(model, vector_store, keyword_store, record_store)

        # --- 1. Corrupt / odd files never abort the scan; good files are searchable ---
        folder_scan.RECENT_FAILURES.clear()
        progress = {}
        def on_progress(done, total, current, failed):
            progress.update(done=done, total=total, failed=failed)
        index_folder(indexer, str(folder), progress=on_progress)
        active = {Path(r.path).name for r in record_store.list_active()}
        assert set(good) <= active, active
        assert search.search("projector broke thursday", top_k=3)[0]["filename"] == "good_one.txt"
        assert search.search("insurance excess", top_k=3)[0]["filename"] == "good_three.txt"
        failures = {Path(f["path"]).name: f["error"] for f in folder_scan.RECENT_FAILURES}
        assert {"broken.pdf", "broken.docx", "broken.xlsx", "broken.pptx"} <= set(failures), failures
        assert all(err for err in failures.values())
        # Item 8 (2026-10-05): a reason the user can read, not "PdfReadError: EOF marker not found".
        assert all(failures[n].startswith("The file is damaged or not a valid ") for n in ("broken.pdf", "broken.docx", "broken.xlsx", "broken.pptx")), failures
        # The CSV formula cells are indexed as plain text — never evaluated.
        hit = search.search('"cmd"', top_k=3)
        assert hit and hit[0]["filename"] == "csv_injection.csv" and "=cmd" in hit[0]["chunk_text"]
        assert progress["failed"] == len({"broken.pdf", "broken.docx", "broken.xlsx", "broken.pptx"}), progress
        print(f"1. {progress['failed']} corrupt files counted and listed with reasons ({', '.join(sorted(failures))}); empty, binary-as-text, 3 MB line and formula-injection CSV handled; good files searchable: OK")

        # --- 2. A file that vanishes between the walk and the read is skipped, not fatal ---
        (folder / "fleeting.txt").write_text("this file will be gone before it is read " * 5)
        real_discover = folder_scan.discover_files
        def discover_then_delete(roots, extensions=None, on_error=None, on_skip=None):
            paths = list(real_discover(roots, extensions, on_error=on_error))
            (folder / "fleeting.txt").unlink()
            return iter(paths)
        folder_scan.discover_files = discover_then_delete
        try:
            index_folder(indexer, str(folder))
        finally:
            folder_scan.discover_files = real_discover
        assert record_store.get_by_path(str(folder / "fleeting.txt")) is None
        assert record_store.get_by_path(str(folder / "good_one.txt")).indexed
        print("2. A file deleted between the walk and the read is skipped; the scan finishes and the rest stays indexed: OK")

        # --- 3. An unreadable file is a counted failure, then indexed once readable ---
        locked = folder / "locked.txt"
        locked.write_text("secret notes about the venue booking for the reunion " * 3)
        set_readable(locked, False)
        folder_scan.RECENT_FAILURES.clear()
        if running_as_root():
            print("3. (running as root — permission check skipped)")
        else:
            index_folder(indexer, str(folder))
            assert any(Path(f["path"]).name == "locked.txt" and f["error"].startswith("Windows refused access") for f in folder_scan.RECENT_FAILURES), folder_scan.RECENT_FAILURES
            assert record_store.get_by_path(str(locked)) is None or not record_store.get_by_path(str(locked)).indexed
            set_readable(locked, True)
            index_folder(indexer, str(folder))
            assert record_store.get_by_path(str(locked)).indexed
            assert search.search("venue booking reunion", top_k=3)[0]["filename"] == "locked.txt"
            print("3. Unreadable file: counted with 'Windows refused access', not fatal; indexed on the next scan once readable: OK")

        # --- 4. Crash mid-index: record saved, chunks never stored → repaired by the next scan, nothing else rescanned ---
        victim = record_store.get_by_path(str(folder / "good_two.txt")) or record_store.get_by_path(str(folder / "good_two.md"))
        indexer.delete_file(victim.file_id)              # simulate: chunks lost…
        record_store.mark_indexed(victim.file_id, False)  # …and the record left "pending", as a crash between the two would
        assert search.search("sourdough cold ferment", top_k=3) == [] or search.search("sourdough cold ferment", top_k=3)[0]["filename"] != "good_two.md"
        assert [r.file_id for r in record_store.list_unfinished()] == [victim.file_id]
        hashed = []
        real_hash = folder_scan.compute_file_hash
        folder_scan.compute_file_hash = lambda p: (hashed.append(p.name), real_hash(p))[1]
        try:
            index_folder(indexer, str(folder))
        finally:
            folder_scan.compute_file_hash = real_hash
        assert "good_two.md" in hashed and "good_one.txt" not in hashed, f"only the unfinished file should be re-read: {hashed}"
        assert record_store.list_unfinished() == []
        assert search.search("sourdough cold ferment", top_k=3)[0]["filename"] == "good_two.md"
        print("4. Crash mid-index repaired by the next scan: only the unfinished file was re-read, the rest skipped on size+mtime: OK")

        # --- 5. An unplugged drive / offline share (2026-10-04). Its folder
        # is missing at scan time: before, every record under it was purged,
        # the folder dropped from the watched list, and a search that missed
        # one file tombstoned it. Now nothing is written while it is away. ---
        usb = workdir / "usb"
        (usb / "sub").mkdir(parents=True)
        (usb / "trip.txt").write_text("Photos from the lighthouse walk on the island, saved on the stick.")
        (usb / "sub" / "deep.txt").write_text("The ferry timetable for the harbour crossing, kept in a sub folder.")
        index_folder(indexer, str(usb))
        trip = str(usb / "trip.txt")
        usb.rename(workdir / "usb_unplugged")
        folder_scan.RECENT_FAILURES.clear()
        index_folder(indexer, str(usb))
        record = record_store.get_by_path(trip)
        assert record is not None and record.indexed and not record.deleted, "files of an unreachable folder must stay indexed"
        assert any(f["path"] == str(usb) and "not reachable" in f["error"] for f in folder_scan.RECENT_FAILURES), folder_scan.RECENT_FAILURES
        assert all(h["path"] != trip for h in search.search("lighthouse walk on the island", top_k=5)), "a file whose folder is away must not be shown"
        assert not record_store.get_by_path(trip).deleted, "search must not tombstone a file whose folder is unreachable"
        (workdir / "usb_unplugged").rename(usb)
        assert search.search("lighthouse walk on the island", top_k=3)[0]["path"] == trip
        # A sub-folder the walk could not list keeps its files too.
        def discover_without_sub(roots, extensions=None, on_error=None, on_skip=None):
            on_error(PermissionError(13, "Access is denied", str(usb / "sub")))
            return iter([p for p in real_discover(roots, extensions) if p.parent != usb / "sub"])
        folder_scan.discover_files = discover_without_sub
        try:
            index_folder(indexer, str(usb))
        finally:
            folder_scan.discover_files = real_discover
        assert not record_store.get_by_path(str(usb / "sub" / "deep.txt")).deleted, "files under an unlistable sub-folder were purged"
        # The watched list keeps a folder that is away, and lists it as missing.
        from app.watch_setup import LiveIndexing, load_watched_folders, save_watched_folders
        config = workdir / "config"
        config.mkdir()
        away = str(workdir / "external_drive")
        save_watched_folders(config, [away])
        assert load_watched_folders(config) == [away]
        live = LiveIndexing(indexer, None, config)
        live.start()
        try:
            live.watch(str(usb))
            assert set(load_watched_folders(config)) == {away, str(usb)}, load_watched_folders(config)
            assert {f["path"]: f["exists"] for f in live.folder_stats()} == {away: False, str(usb): True}
            live.forget(away)
            assert load_watched_folders(config) == [str(usb)]
        finally:
            live.stop()
        # A file deleted while its folder is there is still tombstoned by search.
        (usb / "trip.txt").unlink()
        search.search("lighthouse walk on the island", top_k=5)
        assert record_store.get_by_path(trip).deleted
        print("5. Unreachable folder (unplugged drive): its files stay indexed and hidden, the scan reports it, the watched list keeps it; an unlistable sub-folder keeps its files; a deleted file is still purged: OK")

        # --- 6. A transient failure is retried (2026-10-04). FAILED_HASHES
        # used to remember ANY failure for the life of the process, so a
        # document Word had locked, or a half-written download, was never
        # retried until the app restarted. Bad bytes are still remembered. ---
        busy = folder / "busy.txt"
        busy.write_text("Minutes of the allotment committee: compost delivery on Saturday.")
        real_index_file = indexer.index_file
        def locked_once(path, *a, **k):
            if path.name == "busy.txt":
                raise PermissionError(13, "The process cannot access the file because it is being used by another process", str(path))
            return real_index_file(path, *a, **k)
        indexer.index_file = locked_once
        try:
            index_folder(indexer, str(folder))
        finally:
            indexer.index_file = real_index_file
        assert not any(h in folder_scan.FAILED_HASHES for h in [folder_scan.compute_file_hash(busy)])
        index_folder(indexer, str(folder))
        assert record_store.get_by_path(str(busy)).indexed, "a file that was locked once must be indexed on the next scan"
        assert folder_scan.compute_file_hash(folder / "broken.pdf") in folder_scan.FAILED_HASHES, "a corrupt file is still remembered"
        print("6. A locked file (PermissionError while reading) is retried on the next scan; a corrupt file is still remembered: OK")

        # --- 7. Oversized and hostile files (2026-10-04): refused with a reason, never parsed. ---
        import zipfile
        from docx import Document
        from app.extraction import extractor
        from app.files.discovery import MAX_DOCUMENT_FILE_BYTES
        hostile = workdir / "hostile"
        hostile.mkdir()
        doc = Document()
        doc.add_paragraph("A real letter about the garden hedge and the boundary fence.")
        doc.save(hostile / "real letter.docx")
        doc.save(hostile / "bomb.docx")
        with zipfile.ZipFile(hostile / "bomb.docx", "a", compression=zipfile.ZIP_DEFLATED) as z, z.open("word/media/zeros.bin", "w") as out:
            for _ in range(20):
                out.write(b"\0" * (1024 * 1024))  # 20 MB of zeros packs into ~20 KB: a ratio of ~1000:1
        with open(hostile / "huge.pdf", "wb") as f:
            f.truncate(MAX_DOCUMENT_FILE_BYTES + 1)
        folder_scan.RECENT_FAILURES.clear()
        folder_scan.RECENT_SKIPS.clear()
        index_folder(indexer, str(hostile))
        reasons = {Path(f["path"]).name: f["error"] for f in folder_scan.RECENT_FAILURES}
        skips = {Path(f["path"]).name: f["reason"] for f in folder_scan.RECENT_SKIPS}
        assert record_store.get_by_path(str(hostile / "real letter.docx")).indexed, reasons
        assert "bomb.docx" in reasons and "zip bomb" in reasons["bomb.docx"], reasons
        # Item 6: a file left alone on purpose is a skip, never listed among the failures.
        assert "huge.pdf" not in reasons and "larger than 50 MB" in skips["huge.pdf"] and folder_scan.SKIPPED[str(hostile)]["too_large"] == 1, (reasons, skips)
        real_cap = extractor.MAX_ZIP_UNCOMPRESSED_BYTES
        extractor.MAX_ZIP_UNCOMPRESSED_BYTES = 1000  # the real letter unpacks to more than this
        try:
            extractor.extract_document(hostile / "real letter.docx")
            raise AssertionError("an Office file that unpacks past the cap must be refused")
        except ValueError as e:
            assert "unpack" in str(e), e
        finally:
            extractor.MAX_ZIP_UNCOMPRESSED_BYTES = real_cap
        from PIL import Image
        assert Image.MAX_IMAGE_PIXELS == 210_000_000
        print(f"7. Hostile files: zip bomb refused ({reasons['bomb.docx'][:70]}...), a 50 MB+ PDF skipped with its reason, a real .docx indexed; image pixel cap {Image.MAX_IMAGE_PIXELS:,}: OK")

        # --- 8. OneDrive "online-only" placeholders (2026-10-04): reading one
        # downloads it, so a scan of a OneDrive folder pulled the whole cloud
        # library down. Files with the offline / recall attributes are
        # counted as skipped and never read. ---
        if sys.platform == "win32":
            import ctypes
            cloud = workdir / "OneDrive"
            cloud.mkdir()
            (cloud / "local.txt").write_text("Kept on this PC: the boiler service history and the warranty.")
            placeholder = cloud / "cloud only.txt"
            placeholder.write_text("Only in the cloud: the scanned passport and visa pages.")
            assert ctypes.windll.kernel32.SetFileAttributesW(str(placeholder), 0x1000)  # FILE_ATTRIBUTE_OFFLINE, as OneDrive marks a placeholder
            try:
                count = index_folder(indexer, str(cloud), progress=on_progress)
            finally:
                ctypes.windll.kernel32.SetFileAttributesW(str(placeholder), 0x80)  # FILE_ATTRIBUTE_NORMAL
            assert record_store.get_by_path(str(placeholder)) is None and record_store.get_by_path(str(cloud / "local.txt")).indexed
            assert folder_scan.SKIPPED[str(cloud)]["online_only"] == 1 and progress["failed"] == 0 and count == 1, (folder_scan.SKIPPED, progress, count)
            assert {"path": str(placeholder), "reason": ONLINE_ONLY_REASON} in folder_scan.RECENT_SKIPS, folder_scan.RECENT_SKIPS
            print("8. OneDrive online-only placeholder: skipped and counted ('online_only': 1), never read or failed; the local file is indexed: OK")

        check_review_items(workdir, indexer, record_store, vector_store, on_progress, progress)

        print("\nPhase 11 reliability OK: corrupt files, vanishing files, unreadable files and a crash mid-index are all survived and reported.")
    finally:
        for p in workdir.rglob("*"):
            try:
                os.chmod(p, 0o644)
            except OSError:
                pass
        shutil.rmtree(workdir, ignore_errors=True)


ONLINE_ONLY_REASON = "online-only OneDrive file; make it available offline to index it"


def check_review_items(workdir: Path, indexer, record_store, vector_store, on_progress, progress) -> None:
    """9-14: the 2026-10-05 robustness review. Each failed before its fix."""
    import json
    import sqlite3
    import struct
    import time

    from watchdog.events import FileCreatedEvent

    from app import watch_setup
    from app.files import discovery
    from app.files.identity import build_file_record, compute_file_hash
    from app.files.jobs import Job
    from app.indexing import CHUNKS_TABLE
    from app.watch_setup import LiveIndexing

    # --- 9. Skips are not failures; the watcher reports its skips too; /status keys (item 6) ---
    folder_scan.RECENT_SKIPS.clear()
    real_long = discovery.long_paths_enabled
    discovery.long_paths_enabled = lambda: False
    try:
        heard = []
        deep = Path("C:/" + "deep folder/" * 25 + "notes.txt")
        assert not discovery.is_indexable(deep, on_skip=lambda p, r: heard.append(r)) and heard == [discovery.LONG_PATH_REASON], heard
    finally:
        discovery.long_paths_enabled = real_long
    config = workdir / "config9"
    config.mkdir()
    watched = workdir / "watched9"
    watched.mkdir()
    live = LiveIndexing(indexer, None, config)
    try:
        live.watcher.add_root(str(watched))
        big = watched / "export.csv"
        big.write_bytes(b"a,b\n" * (2 * 1024 * 1024))  # 8 MB, over the 5 MB cap for .csv
        live.watcher._handler.dispatch(FileCreatedEvent(str(big)))
        live.watcher._handler.dispatch(FileCreatedEvent(str(big)))  # a second event: still one entry
        snapshot = live.job_snapshot()
    finally:
        live.stop()
    assert set(snapshot["skipped"]) == {"too_large", "online_only", "path_too_long"}, snapshot["skipped"]
    assert [s["path"] for s in snapshot["recent_skips"]] == [str(big)] and set(snapshot["recent_skips"][0]) == {"path", "reason"}, snapshot["recent_skips"]
    assert "larger than 5 MB" in snapshot["recent_skips"][0]["reason"]
    assert not any("Skipped" in f["error"] for f in snapshot["recent_failures"]), snapshot["recent_failures"]
    print("9. Skips kept apart from failures (job.skipped by kind, job.recent_skips); the watcher reports a too-large file; a long path is a skip with its reason: OK")

    # --- 10. Failure reasons are plain sentences (item 8) ---
    from pypdf.errors import FileNotDecryptedError, PdfStreamError

    try:
        try:
            struct.unpack("<I", b"")
        except struct.error as e:
            raise ValueError(f"old .doc could not be read: error: {e}") from e
    except ValueError as e:
        damaged_doc = e
    folder_scan.RECENT_FAILURES.clear()
    cases = {
        "a.pdf": (PdfStreamError("Stream has ended unexpectedly"), "The file is damaged or not a valid PDF file."),
        "old.doc": (damaged_doc, "The file is damaged or not a valid Word file."),
        "open.docx": (PermissionError(13, "The process cannot access the file", "open.docx"), "Windows refused access (the file may be open in another program)."),
        "huge.tif": (MemoryError(), "Too large to read on this computer."),
        "secret.pdf": (FileNotDecryptedError("File has not been decrypted"), "The file is password-protected."),
        "locked.doc": (ValueError("old .doc could not be read: it is password-protected"), "The file is password-protected."),
        "odd.txt": (ValueError("its text could not be decoded"), "Its text could not be decoded."),
    }
    for name, (error, _) in cases.items():
        folder_scan._note_failure(workdir / name, error)
    got = {Path(f["path"]).name: f.get("reason") for f in folder_scan.RECENT_FAILURES}
    assert got == {n: want for n, (_, want) in cases.items()}, got
    assert all(f["error"] == f["reason"] for f in folder_scan.RECENT_FAILURES)
    print("10. Failure reasons are plain sentences (damaged <type>, password-protected, refused access, too large), the raw error stays in the log: OK")

    # --- 11. A scan that raises: job.state 'failed' with a readable error (item 7) ---
    config = workdir / "config11"
    config.mkdir()
    live = LiveIndexing(indexer, None, config)
    real_index_folder = watch_setup.index_folder

    def broken_scan(*a, **k):
        raise sqlite3.OperationalError("database is locked")
    watch_setup.index_folder = broken_scan
    try:
        live.enqueue(str(workdir))
        live._runner.join(timeout=10)
    finally:
        watch_setup.index_folder = real_index_folder
    job = live.job_snapshot()
    assert job["state"] == "failed" and job["error"].startswith("The scan stopped") and "database is locked" in job["error"] and "OperationalError" not in job["error"], job
    print(f"11. A scan that fails: job.state 'failed', job.error {job['error']!r}: OK")

    # --- 12. A file that crashes the app is skipped until it changes (item 4) ---
    crashy = workdir / "crashy"
    crashy.mkdir()
    poison = crashy / "poison.txt"
    poison.write_text("A spreadsheet of hedgehog sightings by month.")
    (crashy / "fine.txt").write_text("Bicycle bell repair receipts.")
    config = workdir / "config12"
    config.mkdir()
    LiveIndexing(indexer, None, config)  # the app's data folder holds the marker
    marker = getattr(folder_scan, "IN_PROGRESS_FILE", None)
    assert marker is not None and marker.parent == config, marker
    seen_marker, calls = [], []
    real_index_file = indexer.index_file

    def spy(path, *a, **k):
        seen_marker.append(json.loads(marker.read_text())["path"] if marker.exists() else None)
        calls.append(path.name)
        return real_index_file(path, *a, **k)
    indexer.index_file = spy
    try:
        index_folder(indexer, str(crashy))
        assert str(poison) in seen_marker and not marker.exists(), (seen_marker, marker.exists())
        # The app dies while reading poison.txt's new version: the marker stays behind.
        poison.write_text("A spreadsheet of hedgehog sightings by month, now with a native crash.")
        new_hash = compute_file_hash(poison)
        record_store.upsert(build_file_record(poison, file_id=record_store.get_by_path(str(poison)).file_id, file_hash=new_hash, indexed=False))
        # First leftover marker: closing the window mid-file looks the same, so
        # the file is only a suspect and is still read.
        marker.write_text(json.dumps({"path": str(poison), "hash": new_hash, "started_at": time.time()}))
        folder_scan.RECENT_FAILURES.clear()
        folder_scan.POISONED.clear(); folder_scan.SUSPECTS.clear()  # a new process
        live = LiveIndexing(indexer, None, config)
        live.start()
        live.stop()
        assert new_hash not in folder_scan.POISONED and new_hash in folder_scan.SUSPECTS, "one interrupted read must not skip the file"
        # Second leftover marker in a row for the same bytes: now it is skipped.
        marker.write_text(json.dumps({"path": str(poison), "hash": new_hash, "started_at": time.time()}))
        folder_scan.POISONED.clear(); folder_scan.SUSPECTS.clear()  # a new process (suspects reload from disk)
        live = LiveIndexing(indexer, None, config)
        live.start()
        live.stop()
        crash_reason = "IntelliFile closed unexpectedly while reading this file; it is skipped until it changes"
        assert any(f["path"] == str(poison) and f["reason"] == crash_reason for f in folder_scan.RECENT_FAILURES), folder_scan.RECENT_FAILURES
        assert not marker.exists()
        calls.clear()
        index_folder(indexer, str(crashy), progress=on_progress)
        assert "poison.txt" not in calls and progress["failed"] == 1, (calls, progress)
        live._handle_job(Job(kind="index", path=str(poison), file_id=record_store.get_by_path(str(poison)).file_id))
        assert "poison.txt" not in calls, "the watcher must skip it too"
        folder_scan.POISONED.clear(); folder_scan.SUSPECTS.clear()  # and the next launch still remembers it
        live = LiveIndexing(indexer, None, config)
        live.start()
        live.stop()
        index_folder(indexer, str(crashy))
        assert "poison.txt" not in calls, calls
        index_folder(indexer, str(crashy), force=True)  # Re-index = the user asking to try again
        assert "poison.txt" in calls and record_store.get_by_path(str(poison)).indexed, ("Re-index must retry it", calls)
        calls.clear()
        poison.write_text("A spreadsheet of hedgehog sightings, saved again by the user.")
        index_folder(indexer, str(crashy))
        assert "poison.txt" in calls and record_store.get_by_path(str(poison)).indexed, calls
    finally:
        indexer.index_file = real_index_file
    print("12. A file the app stopped reading twice in a row (once = a suspect, still read) is marked failed at the next start and skipped by scans and the watcher until it changes; Re-index retries it: OK")

    # --- 13. A retried live job never doubles the file's chunks (item 15) ---
    once = workdir / "dup13"
    once.mkdir()
    note = once / "note.txt"
    note.write_text("The quokka bicycle repair shop opens at eight.")
    record = build_file_record(note, indexed=False)
    record_store.upsert(record)
    config = workdir / "config13"
    config.mkdir()
    live = LiveIndexing(indexer, None, config)
    real_add = indexer.keyword_store.add_chunks

    def fail_once(chunks):
        indexer.keyword_store.add_chunks = real_add
        raise sqlite3.OperationalError("disk I/O error")
    indexer.keyword_store.add_chunks = fail_once
    try:
        live._handle_job(Job(kind="index", path=str(note), file_id=record.file_id, new=True))  # vectors written, keywords not
        live._handle_job(Job(kind="index", path=str(note), file_id=record.file_id, new=True))  # the retry
    finally:
        indexer.keyword_store.add_chunks = real_add
    rows = vector_store.get_by_file_id(CHUNKS_TABLE, record.file_id)
    assert len(rows) == 1, f"{len(rows)} copies of a one-chunk file after a retried job"
    print("13. A live job retried after a half-written attempt replaces its chunks, never duplicates them: OK")

    # --- 14. A dead network drive never blocks /status (item 5) ---
    config = workdir / "config14"
    config.mkdir()
    live = LiveIndexing(indexer, None, config)
    dead = str(workdir / "dead_share")
    live.unreachable_roots.add(dead)
    real_is_dir = Path.is_dir

    def hanging_is_dir(self):
        if str(self) == dead:
            time.sleep(3)
            return False
        return real_is_dir(self)
    Path.is_dir = hanging_is_dir
    try:
        started = time.perf_counter()
        for _ in range(3):  # three /status polls
            stats = live.folder_stats()
        elapsed = time.perf_counter() - started
    finally:
        time.sleep(3.2)  # let a background probe finish before restoring
        Path.is_dir = real_is_dir
    assert elapsed < 0.5 and {f["path"]: f["exists"] for f in stats}[dead] is False, (elapsed, stats)
    print(f"14. A watched folder on a dead network drive: 3 status polls took {elapsed * 1000:.0f} ms (probe runs in the background): OK")

    # --- 15. Installed software and toolchains are not the user's files ---
    from app.files.discovery import discover_files, under_excluded_directory
    pc = workdir / "pc15"
    for rel in ["Documents/plan.txt", "AppData/Local/app/cache.json", "anaconda3/Lib/os.py", "tools/myenv/pyvenv.cfg",
                "tools/myenv/Lib/mod.py", "DevTools/rust/toolchains/x/lib/rustlib/src/core.rs", "Games/App/unins000.exe",
                "Games/App/readme.txt", "Projects/app/windows/runner/main.cpp", "Projects/app/notes.md"]:
        f = pc / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text("x")
    found = sorted(str(p.relative_to(pc)).replace("\\", "/") for p in discover_files([pc]))
    assert found == ["Documents/plan.txt", "Projects/app/notes.md", "Projects/app/windows/runner/main.cpp"], found
    assert under_excluded_directory(pc / "tools/myenv/Lib/new.py", pc) and under_excluded_directory(pc / "Games/App/new.txt", pc)
    assert not under_excluded_directory(pc / "Documents/new.txt", pc)
    print("15. Installed software is skipped (AppData, anaconda3, a venv by pyvenv.cfg, a Rust toolchain, an installed program); the user's documents and project code are kept, for the scan and the live watcher: OK")


if __name__ == "__main__":
    main()
