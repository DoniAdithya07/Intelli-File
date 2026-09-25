"""Phase 3 integration test: build a real PDF, DOCX, TXT, MD (and since
2026-09-20: CSV, XLSX, PPTX, HTML, code, RTF) file,
extract text from each, chunk the result, and verify page numbers /
headings / overlap all come out correctly. Run with:

    backend/venv/bin/python backend/scripts/prototype_extraction.py

Requires backend/requirements-dev.txt (fpdf2, for generating the test PDF).
"""

import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.chunking import chunk_document  # noqa: E402
from app.extraction.extractor import extract_document  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402
from fpdf import FPDF  # noqa: E402


def make_test_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)

    pdf.add_page()
    pdf.multi_cell(0, 10, "Page one talks about horizontal scaling and load balancing for traffic spikes.")

    pdf.add_page()
    pdf.multi_cell(0, 10, "Page two talks about something unrelated, like baking sourdough bread at home.")

    pdf.output(str(path))


def make_test_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("System Design Notes", level=1)  # -> section
    doc.add_heading("Scaling", level=2)  # -> heading
    doc.add_paragraph("Horizontal scaling allows additional instances to be provisioned as demand increases.")
    doc.add_paragraph("This works well combined with a load balancer distributing traffic across instances.")
    doc.add_heading("Storage", level=2)  # -> heading
    doc.add_paragraph("Embedded databases avoid the need for a separate server process.")
    doc.save(str(path))


def main():
    workdir = Path(tempfile.mkdtemp())

    try:
        # --- TXT ---
        txt_path = workdir / "notes.txt"
        txt_path.write_text("Kubernetes autoscaling adjusts pod count based on CPU usage automatically.")
        txt_blocks = extract_document(txt_path)
        assert len(txt_blocks) == 1, "TXT should be a single block"
        assert txt_blocks[0].page_number is None and txt_blocks[0].heading is None
        assert "Kubernetes" in txt_blocks[0].text
        print("1. TXT extraction: OK")

        # --- Markdown ---
        md_path = workdir / "readme.md"
        md_path.write_text("# Title\n\nSome content about distributed systems.")
        md_blocks = extract_document(md_path)
        assert len(md_blocks) == 1
        assert "distributed systems" in md_blocks[0].text
        print("2. Markdown extraction: OK")

        # --- DOCX with headings/sections ---
        docx_path = workdir / "design.docx"
        make_test_docx(docx_path)
        docx_blocks = extract_document(docx_path)
        assert len(docx_blocks) == 3, f"Expected 3 paragraph blocks, got {len(docx_blocks)}"
        assert docx_blocks[0].section == "System Design Notes" and docx_blocks[0].heading == "Scaling"
        assert docx_blocks[1].heading == "Scaling"
        assert docx_blocks[2].heading == "Storage"
        assert "Embedded databases" in docx_blocks[2].text
        print("3. DOCX extraction (headings/sections tracked correctly): OK")

        # --- 3b. DOCX tables are extracted, in document order, under the
        # current heading (2026-09-21: `doc.paragraphs` skipped every table,
        # so an invoice or timetable indexed as nearly empty) ---
        table_doc = DocxDocument()
        table_doc.add_heading("Invoice", level=1)
        table_doc.add_paragraph("Thanks for your business.")
        table = table_doc.add_table(rows=2, cols=3)
        for r, row in enumerate([["Item", "Qty", "Price"], ["Kumquat crate", "4", "120"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        table_doc.add_paragraph("Paid by bank transfer.")
        table_path = workdir / "invoice.docx"
        table_doc.save(str(table_path))
        table_blocks = extract_document(table_path)
        assert len(table_blocks) == 3, f"expected paragraph, table, paragraph — got {len(table_blocks)} blocks"
        assert "Kumquat crate | 4 | 120" in table_blocks[1].text and "Item | Qty | Price" in table_blocks[1].text
        assert table_blocks[1].section == "Invoice", "table must carry the heading context"
        assert table_blocks[2].text == "Paid by bank transfer.", "content after the table must keep document order"
        print("3b. DOCX tables extracted in document order (rows kept together, header row searchable): OK")

        # --- PDF with page numbers ---
        pdf_path = workdir / "systemdesign.pdf"
        make_test_pdf(pdf_path)
        pdf_blocks = extract_document(pdf_path)
        assert len(pdf_blocks) == 2, f"Expected 2 pages, got {len(pdf_blocks)}"
        assert pdf_blocks[0].page_number == 1 and "horizontal scaling" in pdf_blocks[0].text.lower()
        assert pdf_blocks[1].page_number == 2 and "bread" in pdf_blocks[1].text.lower()
        print("4. PDF extraction (page numbers tracked correctly): OK")

        # --- Chunking: page number carries through from PDF blocks ---
        pdf_chunks = chunk_document(pdf_blocks, chunk_size=20, overlap=5)
        assert len(pdf_chunks) >= 2, "Expected at least one chunk per page with this chunk size"
        assert pdf_chunks[0].page_number == 1
        assert any(c.page_number == 2 for c in pdf_chunks), "A later chunk should land on page 2"
        print("5. Chunking preserves page_number: OK")

        # 2026-09-20: at the DEFAULT chunk size a short 2-page PDF used to
        # collapse into one chunk labelled page 1, so any page-2 hit said
        # "Found on page 1". Chunks must never span a page boundary.
        default_chunks = chunk_document(pdf_blocks)
        assert [c.page_number for c in default_chunks] == [1, 2], [c.page_number for c in default_chunks]
        assert "bread" in default_chunks[1].content and "bread" not in default_chunks[0].content
        assert default_chunks[1].start_offset > default_chunks[0].end_offset, "offsets must stay document-global"
        print("5b. Chunks never cross a page boundary at the default chunk size: OK")

        # --- Chunking: heading/section carries through from DOCX blocks ---
        docx_chunks = chunk_document(docx_blocks, chunk_size=15, overlap=3)
        assert any(c.heading == "Storage" for c in docx_chunks), "Expected a chunk tagged with the Storage heading"
        print("6. Chunking preserves heading/section: OK")

        # --- Chunking: overlap actually overlaps ---
        long_text = " ".join(f"word{i}" for i in range(100))
        from app.extraction.blocks import ExtractedBlock

        chunks = chunk_document([ExtractedBlock(text=long_text)], chunk_size=20, overlap=5)
        assert len(chunks) >= 2
        first_words = set(chunks[0].content.split())
        second_words = set(chunks[1].content.split())
        overlap_words = first_words & second_words
        assert len(overlap_words) == 5, f"Expected exactly 5 overlapping words, got {len(overlap_words)}"
        print(f"7. Sliding-window overlap: OK ({len(overlap_words)} words shared between consecutive chunks, as configured)")

        # --- Chunking: whole document covered, no gaps ---
        all_content = " ".join(c.content for c in chunks)
        for i in range(100):
            assert f"word{i}" in all_content, f"word{i} missing from any chunk — chunking left a gap"
        print("8. Full document coverage (no gaps between chunks): OK")

        # --- Edge case: overlap >= chunk_size must raise, not silently misbehave ---
        try:
            chunk_document([ExtractedBlock(text="a b c")], chunk_size=5, overlap=5)
            raise AssertionError("Expected ValueError when overlap >= chunk_size")
        except ValueError:
            print("9. Invalid overlap/chunk_size correctly rejected: OK")

        # --- Edge case: empty file produces zero chunks, not a crash ---
        empty_path = workdir / "empty.txt"
        empty_path.write_text("")
        empty_blocks = extract_document(empty_path)
        assert empty_blocks == []
        assert chunk_document(empty_blocks) == []
        print("10. Empty file handled gracefully: OK")

        # --- 2026-09-20: spreadsheets, slides, HTML, code, RTF ---
        csv_path = workdir / "invoices.csv"
        csv_path.write_text("customer,amount,status\nAlice Johnson,1200,paid\nBob Stone,450,overdue\n")
        csv_blocks = extract_document(csv_path)
        assert len(csv_blocks) == 1 and "customer amount status" in csv_blocks[0].text
        assert "Alice Johnson 1200 paid" in csv_blocks[0].text, "row cells must stay together on one line"
        semi_path = workdir / "euro.csv"
        semi_path.write_text("name;city\nMarta;Lisbon\n")
        assert "Marta Lisbon" in extract_document(semi_path)[0].text, "semicolon delimiter should be sniffed"
        print("11. CSV extraction (header kept, rows intact, delimiter sniffed): OK")

        from openpyxl import Workbook

        xlsx_path = workdir / "budget.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Q1"
        ws.append(["item", "cost"])
        ws.append(["rent", 900])
        ws.append(["total", "=SUM(B2:B2)"])
        wb.create_sheet("Notes").append(["Remember to renew the gym membership"])
        wb.save(str(xlsx_path))
        xlsx_blocks = extract_document(xlsx_path)
        assert [b.section for b in xlsx_blocks] == ["Q1", "Notes"], xlsx_blocks
        assert "rent 900" in xlsx_blocks[0].text
        assert "SUM(" not in xlsx_blocks[0].text, "formulas must never be surfaced or evaluated"
        assert "gym membership" in xlsx_blocks[1].text
        print("12. XLSX extraction (one block per sheet, formulas not evaluated): OK")

        from pptx import Presentation
        from pptx.util import Inches

        pptx_path = workdir / "pitch.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Market Size"
        slide.placeholders[1].text = "Ten million potential users in Europe"
        slide.notes_slide.notes_text_frame.text = "Mention the survey from March"
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        slide2.shapes.title.text = "Roadmap"
        slide2.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Launch in autumn"
        prs.save(str(pptx_path))
        pptx_blocks = extract_document(pptx_path)
        assert len(pptx_blocks) == 2
        assert pptx_blocks[0].page_number == 1 and pptx_blocks[0].heading == "Market Size"
        assert "potential users" in pptx_blocks[0].text and "survey from March" in pptx_blocks[0].text
        assert pptx_blocks[1].page_number == 2 and "Launch in autumn" in pptx_blocks[1].text
        print("13. PPTX extraction (slide number as page, title as heading, notes included): OK")

        html_path = workdir / "page.html"
        html_path.write_text("<html><head><title>Recipe</title><style>p{color:red}</style></head>"
                             "<body><script>alert('x')</script><h1>Sourdough</h1><p>Mix flour &amp; water.</p></body></html>")
        html_blocks = extract_document(html_path)
        assert len(html_blocks) == 1
        assert "Sourdough" in html_blocks[0].text and "Mix flour & water." in html_blocks[0].text
        assert "alert" not in html_blocks[0].text and "color" not in html_blocks[0].text, "scripts/styles must be dropped"
        print("14. HTML extraction (visible text only, scripts/styles dropped): OK")

        py_path = workdir / "scaler.py"
        py_path.write_text("def autoscale(cpu):\n    # scale pods when cpu is high\n    return cpu > 80\n")
        py_blocks = extract_document(py_path)
        assert len(py_blocks) == 1 and "scale pods when cpu is high" in py_blocks[0].text
        rtf_path = workdir / "memo.rtf"
        rtf_path.write_text(r"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0\fs24 Quarterly memo about hiring.\par Second paragraph.}")
        rtf_blocks = extract_document(rtf_path)
        assert "Quarterly memo about hiring." in rtf_blocks[0].text and "Second paragraph." in rtf_blocks[0].text
        assert "\\rtf" not in rtf_blocks[0].text and "fonttbl" not in rtf_blocks[0].text, rtf_blocks[0].text
        print("15. Code + RTF extraction (plain text, RTF control words stripped): OK")

        from app.files.discovery import discover_files, is_indexable

        (workdir / "package-lock.json").write_text("{}")
        (workdir / "app.min.js").write_text("x")
        big = workdir / "dump.csv"
        with big.open("wb") as fh:
            fh.truncate(6 * 1024 * 1024)
        found = {p.name for p in discover_files([workdir])}
        assert {"invoices.csv", "budget.xlsx", "pitch.pptx", "page.html", "scaler.py", "memo.rtf"} <= found, found
        assert not {"package-lock.json", "app.min.js", "dump.csv"} & found, found
        assert not is_indexable(big), "6 MB csv must be over the size cap"
        print("16. Discovery: new types found; lockfiles, minified bundles and >5 MB data files skipped: OK")

        print("\nPhase 3 extraction + chunking OK: PDF/DOCX/TXT/MD + CSV/XLSX/PPTX/HTML/code/RTF extraction, page/heading tracking, sliding-window overlap, and edge cases all work as expected.")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


if __name__ == "__main__":
    main()
