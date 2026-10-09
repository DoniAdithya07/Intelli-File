"""End-to-end feature check: does IntelliFile identify files, photos and
videos correctly? Run it before every release or push.

It builds a small test folder of known content, starts a real backend on
it (the source backend, or the packaged one with --packaged), and asks the
same HTTP API the desktop app uses. Every check prints what it expected,
what came back, and PASS or FAIL.

    file identification   meaning, exact words, misspelled file names, every
                          document format, type filters, the preview passage
    photo identification  a described picture finds it; nonsense finds nothing;
                          words inside a screenshot are found (Windows OCR)
    video identification  a described scene finds the video and the right
                          moment; the frame at that moment can be shown
    audio identification  words spoken in a recording find it (when Windows'
                          speech engine is available to make the recording)
    safety                0 network attempts (offline guard on) and the test
                          files are byte-identical afterwards (read-only)

Your own index is never touched: the backend runs on a throwaway data
folder with its own API token, and everything is deleted afterwards.

Run (from backend/):
    venv\\Scripts\\python.exe scripts\\check_features.py
    venv\\Scripts\\python.exe scripts\\check_features.py --packaged ..\\desktop\\src-tauri\\target\\release\\IntelliFile
Exit code 0 means every check passed. A JSON report is written with --report.
When a check fails, the evidence for its cause is collected at once and
printed under it: whether the file is in the index, the text IntelliFile
extracted from it, the results with their scores, the search route, the
models, and the backend's error lines.

    --fault corrupt-pdf | blank-screenshot | silent-recording
        breaks one test file on purpose; the run must then fail, and the
        diagnosis must point at the cause (a self-test of this checker).
"""

import argparse
import hashlib
import json
import os
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import requests
from PIL import Image, ImageDraw, ImageFont

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

RESULTS: list[dict] = []


def check(group: str, name: str, passed: bool, evidence: str, diagnose=None) -> None:
    """Records one check. When it fails and `diagnose` is given, the
    evidence for finding the cause is collected at once and kept with it."""
    result = {"group": group, "check": name, "result": "PASS" if passed else "FAIL", "evidence": evidence}
    print(f"  [{result['result']}] {name}: {evidence}", flush=True)
    if not passed and diagnose is not None:
        try:
            result["diagnosis"] = diagnose()
        except Exception as e:  # a diagnosis must never hide the failure itself
            result["diagnosis"] = {"diagnosis_error": f"{type(e).__name__}: {e}"}
        for key, value in result["diagnosis"].items():
            print(f"        {key}: {json.dumps(value, ensure_ascii=False)[:300]}", flush=True)
    RESULTS.append(result)


def skip(group: str, name: str, reason: str) -> None:
    RESULTS.append({"group": group, "check": name, "result": "SKIP", "evidence": reason})
    print(f"  [SKIP] {name}: {reason}", flush=True)


# ---------------------------------------------------------------- test folder

DOCUMENTS = {
    "scaling_notes.md": "# System design notes\n\nHorizontal scaling allows additional server instances to be provisioned when traffic demand increases, distributed by a load balancer.\n",
    "march invoice.txt": "Invoice 2025-03: consulting services, 12 hours. Payment terms net 30, so it is due on 14 April 2025; the amount is 1,440 euros.\n",
    "sourdough.md": "Sourdough bread: mix flour, water, salt and starter, let it rise overnight, bake at high heat.\n",
}
# One unusual phrase per format, so a hit can only come from reading that file.
FORMAT_PHRASES = {
    "turbine manual.pdf": "zephyr turbine maintenance schedule",
    "cake recipe.docx": "lemon drizzle cake with poppy seeds",
    "roadmap.pptx": "quarterly roadmap milestones for the aquarium project",
    "garden log.csv": "hydroponic basil yield",
}
SHAPES = {"red circle.png": ("red", "circle"), "blue square.png": ("blue", "square"), "green triangle.heic": ("green", "triangle")}
# At least 5 words: IntelliFile indexes an image's text only from OCR_MIN_WORDS
# (app/indexing/indexer.py) words, so a photo with one stray sign stays a photo.
SCREENSHOT_TEXT = "Payment received for invoice 4471 from Northwind Traders"
SPOKEN_TEXT = "The quarterly planning meeting has moved to Thursday afternoon."


def draw_shape(path: Path, colour: str, shape: str) -> None:
    img = Image.new("RGB", (640, 480), (245, 245, 240))
    d = ImageDraw.Draw(img)
    if shape == "circle":
        d.ellipse((170, 90, 470, 390), fill=colour)
    elif shape == "square":
        d.rectangle((170, 90, 470, 390), fill=colour)
    else:
        d.polygon([(320, 80), (120, 400), (520, 400)], fill=colour)
    if path.suffix == ".heic":  # an iPhone photo: proves HEIC decoding (pi-heif) made it into the build
        import pillow_heif  # requirements-dev.txt; writes HEIC, the app only reads it
        pillow_heif.from_pillow(img).save(str(path), quality=90)
    else:
        img.save(path)


# --fault breaks one test file on purpose, to prove that the check fails and
# that its diagnosis names the cause (a self-test of the checker and of the
# feature-checker agent). Each is a real-world problem, not an app change.
FAULTS = {
    "corrupt-pdf": "the PDF is replaced by bytes that are not a PDF (a damaged download)",
    "blank-screenshot": "the screenshot has no text in it (nothing for OCR to read)",
    "silent-recording": "the recording is silence (nothing for speech recognition to hear)",
}


def make_folder(root: Path, fault: str | None = None) -> dict:
    """Writes the test folder; returns what was made, for the checks."""
    docs, photos, videos, audio = root / "documents", root / "photos", root / "videos", root / "audio"
    for d in (docs, photos, videos):
        d.mkdir(parents=True)
    for name, text in DOCUMENTS.items():
        (docs / name).write_text(text, encoding="utf-8")

    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.multi_cell(0, 10, f"Plant operations. The {FORMAT_PHRASES['turbine manual.pdf']} is reviewed every spring.")
    pdf.output(str(docs / "turbine manual.pdf"))
    if fault == "corrupt-pdf":
        (docs / "turbine manual.pdf").write_bytes(b"%PDF-1.4\n" + os.urandom(4000))
    import docx
    document = docx.Document()
    document.add_paragraph(f"Weekend baking: {FORMAT_PHRASES['cake recipe.docx']}, served cold.")
    document.save(str(docs / "cake recipe.docx"))
    import pptx
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Planning"
    slide.placeholders[1].text = FORMAT_PHRASES["roadmap.pptx"]
    deck.save(str(docs / "roadmap.pptx"))
    (docs / "garden log.csv").write_text(f"week,crop,note\n1,basil,{FORMAT_PHRASES['garden log.csv']} doubled\n", encoding="utf-8")

    for name, (colour, shape) in SHAPES.items():
        draw_shape(photos / name, colour, shape)
    shot = Image.new("RGB", (1900, 300), "white")
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 64)
    except OSError:
        font = ImageFont.load_default()
    if fault != "blank-screenshot":
        ImageDraw.Draw(shot).text((40, 110), SCREENSHOT_TEXT, fill="black", font=font)
    shot.save(photos / "receipt screenshot.png")

    import av
    frames = []
    for colour, shape in (("red", "circle"), ("blue", "square")):
        img = Image.new("RGB", (640, 360), "white")
        d = ImageDraw.Draw(img)
        (d.ellipse if shape == "circle" else d.rectangle)([200, 40, 440, 320], fill=colour)
        frames.append(img)
    with av.open(str(videos / "shapes clip.mp4"), mode="w") as container:  # red circle for 3 s, then a blue square for 3 s
        stream = container.add_stream("libx264", rate=2)
        stream.width, stream.height, stream.pix_fmt = 640, 360, "yuv420p"
        stream.options = {"g": "1"}
        for img in frames:
            for _ in range(6):
                for packet in stream.encode(av.VideoFrame.from_image(img)):
                    container.mux(packet)
        for packet in stream.encode():
            container.mux(packet)

    made_audio = False
    if fault == "silent-recording":  # 4 s of silence in the format the app records (16 kHz mono 16-bit)
        import wave
        audio.mkdir()
        with wave.open(str(audio / "meeting note.wav"), "wb") as w:
            w.setnchannels(1)
            w.setsampwidth(2)
            w.setframerate(16000)
            w.writeframes(b"\x00\x00" * 16000 * 4)
        made_audio = True
    try:
        import speech_synth
        if not made_audio and speech_synth.available():
            audio.mkdir()
            speech_synth.synthesize_wav(SPOKEN_TEXT, audio / "meeting note.wav")
            made_audio = (audio / "meeting note.wav").stat().st_size > 10_000
    except Exception as e:  # no speech engine: the audio check is skipped, not failed
        print(f"  (no test recording: {type(e).__name__}: {e})")
    n_docs = len(DOCUMENTS) + len(FORMAT_PHRASES)
    return {"documents": n_docs, "photos": len(SHAPES) + 1, "videos": 1, "audio": 1 if made_audio else 0}


def fingerprint(root: Path) -> dict[str, str]:
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(root.rglob("*")) if p.is_file()}


# ---------------------------------------------------------------- backend

def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class Backend:
    def __init__(self, packaged: Path | None, data_dir: Path):
        self.port, self.token = free_port(), secrets.token_hex(16)
        env = {**os.environ, "INTELLIFILE_OFFLINE_GUARD": "1", "INTELLIFILE_APP_DATA_DIR": str(data_dir),
               "INTELLIFILE_API_TOKEN": self.token, "PYTHONIOENCODING": "utf-8"}
        env.pop("INTELLIFILE_PARENT_PID", None)
        if packaged:
            exe = packaged / "backend" / "intellifile-backend.exe"
            cmd, cwd = [str(exe), "--port", str(self.port)], exe.parent
        else:
            cmd, cwd = [sys.executable, "-m", "uvicorn", "app.main:app", "--port", str(self.port)], BACKEND
        self.log_path = data_dir.parent / "backend.log"
        self.log = open(self.log_path, "w", encoding="utf-8")
        self.proc = subprocess.Popen(cmd, cwd=cwd, env=env, stdout=self.log, stderr=subprocess.STDOUT)
        self.url = f"http://127.0.0.1:{self.port}"

    def get(self, endpoint: str, **params):
        return requests.get(self.url + endpoint, params=params, headers={"X-IntelliFile-Token": self.token}, timeout=120)

    def post(self, endpoint: str, body: dict):
        return requests.post(self.url + endpoint, json=body, headers={"X-IntelliFile-Token": self.token}, timeout=120)

    def wait_healthy(self, timeout: float = 180) -> float:
        t0 = time.time()
        while time.time() - t0 < timeout:
            if self.proc.poll() is not None:
                raise RuntimeError(f"backend exited with code {self.proc.returncode}")
            try:
                if requests.get(self.url + "/health", timeout=2).ok:
                    return time.time() - t0
            except requests.RequestException:
                pass
            time.sleep(0.5)
        raise RuntimeError(f"backend not healthy after {timeout:.0f} s")

    def stop(self) -> None:
        self.proc.terminate()
        try:
            self.proc.wait(timeout=20)
        except subprocess.TimeoutExpired:
            self.proc.kill()
        self.log.close()


def wait_indexed(b: Backend, expected_files: int, timeout: float = 300) -> dict:
    """Until every test file is either indexed or recorded as failed: a file
    that cannot be read (a damaged PDF) never joins the index, and waiting
    for it would only end at the timeout."""
    t0, status = time.time(), {}
    while time.time() - t0 < timeout:
        status = b.get("/status").json()
        job = status["job"]
        settled = status["totals"]["files"] + len(job.get("recent_failures") or [])
        if job["state"] != "running" and not job["queued"] and settled >= expected_files:
            return status
        time.sleep(1)
    print(f"  (indexing did not settle within {timeout:.0f} s)", flush=True)
    return status


# ---------------------------------------------------------------- checks

def top_names(results: list[dict], n: int = 3) -> list[str]:
    return [r["filename"] for r in results[:n]]


def ranking(results: list[dict], n: int = 5) -> list[dict]:
    """What the engine returned, with the numbers behind the order."""
    keep = ("filename", "confidence", "score", "keyword_score", "semantic_score", "reranker_score", "timestamp_offset_seconds")
    out = []
    for r in results[:n]:
        row = {k: (round(v, 4) if isinstance(v, float) else v) for k, v in r.items() if k in keep and v is not None}
        if r.get("why"):
            row["why"] = r["why"][:2]
        out.append(row)
    return out


def log_errors(log_path: Path, limit: int = 8) -> list[str]:
    """Error lines the backend wrote while the check ran."""
    try:
        lines = log_path.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []
    hits = [l.strip() for l in lines if any(w in l for w in ("ERROR", "Traceback", "Error:", "Exception", "WARNING"))
            and "Offline guard ON" not in l]  # the start-up notice of the guard is not an error
    return hits[-limit:]


def run_checks(b: Backend, made: dict, status: dict) -> None:
    totals = status.get("totals", {})
    models = status.get("models", {})
    failures = status.get("job", {}).get("recent_failures") or []
    last_route: dict = {}

    def search(q: str, **extra) -> list[dict]:
        body = b.get("/search", q=q, top_k=10, remember="false", **extra).json()
        last_route.clear()
        last_route.update(body.get("route") or {})
        return body.get("results", [])

    # Every indexed file by name, so a failure can say whether the expected
    # file is in the index at all and what text IntelliFile got out of it.
    catalogue: dict[str, dict] = {}
    for kind in ("document", "image", "video", "audio"):
        for r in b.get("/search", q=f"type:{kind}", top_k=100, remember="false").json().get("results", []):
            catalogue[r["filename"]] = r

    def file_facts(name: str) -> dict:
        hit = catalogue.get(name)
        failed = [f["error"][:160] for f in failures if Path(f["path"]).name == name]
        if hit is None:
            return {"in_index": False, "indexing_error": failed[0] if failed else "none recorded"}
        text = b.get("/file-text", file_id=hit["file_id"]).json().get("text", "")
        facts = {"in_index": True, "text_extracted_chars": len(text), "text_start": text[:120]}
        if failed:
            facts["indexing_error"] = failed[0]
        return facts

    def search_diag(expected: str, results: list[dict]):
        return lambda: {"expected_file": {expected: file_facts(expected)},
                        "results": ranking(results),
                        "route": {k: last_route.get(k) for k in ("tier", "requested_tier", "escalated", "reason")},
                        "backend_errors": log_errors(b.log_path)}

    print("\nIndexing")
    for kind in ("documents", "photos", "videos", "audio"):
        if kind == "audio" and not made["audio"]:
            continue
        check("indexing", f"{kind} indexed", totals.get(kind) == made[kind], f"expected {made[kind]}, index reports {totals.get(kind)}",
              lambda: {"files_that_failed": [{"file": Path(f["path"]).name, "error": f["error"][:200]} for f in failures],
                       "models": {k: (v or {}).get("name") if isinstance(v, dict) else v for k, v in models.items()},
                       "backend_errors": log_errors(b.log_path)})
    check("indexing", "no file failed to index", not failures, f"{len(failures)} failures" + (f": {Path(failures[0]['path']).name} ({failures[0]['error'][:80]})" if failures else ""),
          lambda: {"failures": [{"file": Path(f["path"]).name, "error": f["error"][:300]} for f in failures], "backend_errors": log_errors(b.log_path)})

    print("\nFile identification")
    r = search("how do we add capacity when lots of visitors arrive")
    check("files", "found by meaning (no shared words)", bool(r) and r[0]["filename"] == "scaling_notes.md", f"top results {top_names(r)}",
          search_diag("scaling_notes.md", r))
    if r:
        p = b.get("/passages", file_id=r[0]["file_id"], chunk_id=r[0].get("chunk_id") or "").json()
        text = (p.get("match") or {}).get("text", "")
        check("files", "preview shows the matching passage", "load balancer" in text, f"passage: {text[:70]!r}",
              lambda: {"passages_response": {k: p.get(k) for k in ("error", "total")}, "file": {r[0]["filename"]: file_facts(r[0]["filename"])}})
    r = search('"payment terms net 30"')
    check("files", "found by exact words", bool(r) and r[0]["filename"] == "march invoice.txt", f"top results {top_names(r)}",
          search_diag("march invoice.txt", r))
    r = search("march invoise")
    check("files", "found by a misspelled file name", bool(r) and r[0]["filename"] == "march invoice.txt", f"top results {top_names(r)}",
          search_diag("march invoice.txt", r))
    for name, phrase in FORMAT_PHRASES.items():
        r = search(phrase)
        fmt = name.rsplit(".", 1)[1].upper()
        check("files", f"{fmt} read by its content", bool(r) and r[0]["filename"] == name, f"'{phrase}' -> {top_names(r)}",
              search_diag(name, r))
    r = search("type:pdf turbine maintenance")
    check("files", "type filter keeps only that type", bool(r) and all(x["filename"].endswith(".pdf") for x in r), f"type:pdf -> {top_names(r, 5)}",
          search_diag("turbine manual.pdf", r))
    r = search("zqxv wlorbt fnarp")
    strong = [x["filename"] for x in r if x.get("confidence") == "strong"]
    check("files", "nonsense finds no confident match", not strong, f"strong matches: {strong or 'none'}",
          lambda: {"results": ranking(r), "route": dict(last_route)})

    print("\nPhoto identification")
    if models.get("photos") is None:
        skip("photos", "photo search", "the photo model (CLIP) is not installed")
    else:
        for name, (colour, shape) in SHAPES.items():
            q = f"a {colour} {shape}"
            res = b.get("/search-visual", q=q, kind="photo").json().get("results", [])
            ok = bool(res) and res[0]["filename"] == name and res[0].get("confidence") == "strong"
            check("photos", f"'{q}' finds {name}", ok, f"top {[(x['filename'], x.get('confidence')) for x in res[:3]]}",
                  lambda res=res, name=name: {"expected_file": {name: {"in_index": name in catalogue}}, "results": ranking(res),
                                              "photo_model": (models.get("photos") or {}).get("name"), "backend_errors": log_errors(b.log_path)})
        v = b.get("/search-visual", q="zqxv wlorbt", kind="photo").json()
        check("photos", "nonsense description finds nothing", not v.get("results"), f"results {len(v.get('results', []))}, unrecognized {v.get('unrecognized')}",
              lambda: {"results": ranking(v.get("results", [])), "unrecognized": v.get("unrecognized")})
    if not models.get("ocr"):
        skip("photos", "words inside a screenshot (OCR)", "Windows OCR is not available on this PC")
    else:
        r = search("payment received northwind traders")
        check("photos", "words inside a screenshot are found (OCR)", any(x["filename"] == "receipt screenshot.png" for x in r[:3]), f"top results {top_names(r)}",
              lambda: {"screenshot": file_facts("receipt screenshot.png"), "ocr": models.get("ocr"),
                       "rule": "an image's text is indexed only when OCR reads at least 5 words (OCR_MIN_WORDS in app/indexing/indexer.py)",
                       "results": ranking(r)})

    print("\nVideo identification")
    if models.get("photos") is None:
        skip("videos", "video search", "the photo model (CLIP) is not installed")
    else:
        moments = {}
        for q, window, label in (("a red circle", (0, 3), "first scene (0-3 s)"), ("a blue square", (3, 6), "second scene (3-6 s)")):
            res = b.get("/search-visual", q=q, kind="video").json().get("results", [])
            hit = next((x for x in res if x["filename"] == "shapes clip.mp4"), None)
            t = hit.get("timestamp_offset_seconds") if hit else None
            moments[q] = t
            vdiag = (lambda res=res, hit=hit: {"video_in_index": "shapes clip.mp4" in catalogue, "results": ranking(res),
                                               "moments_found": (hit or {}).get("moments"), "backend_errors": log_errors(b.log_path)})
            check("videos", f"'{q}' finds the video", hit is not None and res[0]["filename"] == "shapes clip.mp4", f"top {top_names(res)}", vdiag)
            check("videos", f"'{q}' points to the {label}", t is not None and window[0] <= t < window[1], f"matched moment {t} s", vdiag)
        video = next((x for x in b.get("/search-visual", q="a blue square", kind="video").json().get("results", []) if x["filename"] == "shapes clip.mp4"), None)
        if video and video.get("path"):
            th = b.get("/thumbnail", path=video["path"], size=160, t=moments.get("a blue square") or 0)
            check("videos", "frame at the matched moment can be shown", th.ok and th.headers.get("content-type", "").startswith("image/"), f"HTTP {th.status_code}, {th.headers.get('content-type')}, {len(th.content)} bytes",
                  lambda: {"response": th.text[:300] if not th.headers.get("content-type", "").startswith("image/") else "image", "backend_errors": log_errors(b.log_path)})

    print("\nAudio identification")
    if not made["audio"]:
        skip("audio", "spoken words in a recording", "no speech engine to make a test recording")
    elif models.get("speech") is None:
        skip("audio", "spoken words in a recording", "the speech model (Whisper) is not installed")
    else:
        r = search("quarterly planning meeting thursday")
        check("audio", "spoken words find the recording", any(x["filename"] == "meeting note.wav" for x in r[:3]), f"top results {top_names(r)}",
              lambda: {"recording": file_facts("meeting note.wav"), "speech_model": (models.get("speech") or {}).get("name"),
                       "expected_words": SPOKEN_TEXT, "results": ranking(r)})


def main() -> int:
    sys.stdout.reconfigure(line_buffering=True)  # progress reaches a log file as it happens
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--packaged", type=Path, help="the unzipped IntelliFile folder (holds backend/intellifile-backend.exe)")
    parser.add_argument("--report", type=Path, help="write the results as JSON here")
    parser.add_argument("--keep", action="store_true", help="keep the temporary test folder for inspection")
    parser.add_argument("--fault", choices=sorted(FAULTS), help="break one test file on purpose, to test that the failure and its cause are reported")
    args = parser.parse_args()

    work = BACKEND.parent / ".local" / "tmp"  # inside the project, not the system temp folder
    work.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix="intellifile_feature_check_", dir=work))
    files, data = tmp / "test files", tmp / "data"
    data.mkdir()
    print(f"IntelliFile feature check ({'packaged: ' + str(args.packaged) if args.packaged else 'source backend'})")
    print(f"Test folder: {files}")
    if args.fault:
        print(f"Fault injected on purpose: {args.fault}: {FAULTS[args.fault]}")
    made = make_folder(files, args.fault)
    before = fingerprint(files)
    backend = Backend(args.packaged, data)
    try:
        started = backend.wait_healthy()
        print(f"Backend ready in {started:.1f} s on port {backend.port}")
        a = backend.post("/access", {"mode": "limited"}).json()
        q = backend.post("/index-folder", {"folder": str(files)}).json()
        if a.get("error") or q.get("error"):
            check("indexing", "folder accepted for indexing", False, f"access: {a.get('error')}, index: {q.get('error')}")
        t0 = time.time()
        status = wait_indexed(backend, sum(made.values()))
        print(f"Indexed in {time.time() - t0:.1f} s")
        run_checks(backend, made, status)
        print("\nSafety")
        guard = backend.get("/offline-guard").json()
        check("safety", "offline guard on, 0 network attempts", guard.get("enabled") is True and guard.get("blocked_attempts") == 0, f"enabled={guard.get('enabled')}, blocked attempts={guard.get('blocked_attempts')}",
              lambda: {"blocked_attempts": guard.get("blocked", [])[:10], "hint": "each entry names what tried to connect; see app/offline_guard.py"})
        after = fingerprint(files)
        changed = [k for k in before if before[k] != after.get(k)]
        check("safety", "test files unchanged (read-only)", not changed and set(before) == set(after), f"{len(before)} files, {len(changed)} changed",
              lambda: {"changed": changed, "missing": sorted(set(before) - set(after)), "added": sorted(set(after) - set(before))})
    except Exception as e:
        check("run", "backend started and answered", False, f"{type(e).__name__}: {e}",
              lambda: {"backend_log_tail": backend.log_path.read_text(encoding="utf-8", errors="replace").splitlines()[-15:] if backend.log_path.exists() else []})
    finally:
        backend.stop()
        if not args.keep:
            # Windows can hold the backend's files for a moment after it exits.
            for _ in range(10):
                shutil.rmtree(tmp, ignore_errors=True)
                if not tmp.exists():
                    break
                time.sleep(1)

    passed = sum(r["result"] == "PASS" for r in RESULTS)
    failed = [r for r in RESULTS if r["result"] == "FAIL"]
    skipped = sum(r["result"] == "SKIP" for r in RESULTS)
    print(f"\n{passed} passed, {len(failed)} failed, {skipped} skipped")
    for r in failed:
        print(f"  FAILED: {r['group']}: {r['check']} ({r['evidence']})")
    if args.report:
        args.report.write_text(json.dumps({"target": str(args.packaged) if args.packaged else "source", "fault": args.fault,
                                           "passed": passed, "failed": len(failed), "skipped": skipped, "results": RESULTS}, indent=2), encoding="utf-8")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
