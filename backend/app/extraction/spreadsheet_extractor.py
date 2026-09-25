"""CSV/TSV and Excel (.xlsx/.xlsm) extraction. Added 2026-09-20.

A spreadsheet becomes one block per sheet: the header row first, then
each row's cells joined by spaces, rows separated by newlines. Keeping
the header row means column names are searchable ("invoice total", "due
date"), and keeping rows intact means a row's values stay together in
one chunk, so "Alice 2024 paid" finds the row about Alice.

Formulas are never evaluated (Phase 12: file contents are untrusted
input) — openpyxl's data_only=True returns the values Excel last cached.
"""

import csv
from pathlib import Path

from .blocks import ExtractedBlock

# A data dump with 100k rows is not a document a person searches by
# meaning; the first rows carry the schema and enough content to find
# the file. Beyond this the rest is dropped (the file is still indexed).
MAX_SPREADSHEET_ROWS = 2000
MAX_CELL_CHARS = 500


def _clean_cell(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text[:MAX_CELL_CHARS]


def _rows_to_text(rows) -> str:
    lines = []
    for i, row in enumerate(rows):
        if i >= MAX_SPREADSHEET_ROWS:
            break
        cells = [_clean_cell(c) for c in row]
        line = " ".join(c for c in cells if c)
        if line:
            lines.append(line)
    return "\n".join(lines)


def extract_delimited(path: Path) -> list[ExtractedBlock]:
    """CSV / TSV via the stdlib. The delimiter is sniffed from the first
    few KB so European semicolon-CSVs work too; .tsv forces a tab."""
    try:
        raw = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        raw = path.read_text(encoding="latin-1")
    if not raw.strip():
        return []

    if path.suffix.lower() == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","

    reader = csv.reader(raw.splitlines(), delimiter=delimiter)
    text = _rows_to_text(reader)
    return [ExtractedBlock(text=text)] if text else []


def extract_xlsx(path: Path) -> list[ExtractedBlock]:
    """One block per non-empty sheet; the sheet name becomes the block's
    section so a hit can say which tab it came from."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            text = _rows_to_text(sheet.iter_rows(values_only=True))
            if text:
                blocks.append(ExtractedBlock(text=text, section=sheet.title))
    finally:
        workbook.close()
    return blocks
