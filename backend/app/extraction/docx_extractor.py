from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from .blocks import ExtractedBlock


def _body_items(doc):
    """Paragraphs and tables in document order. `doc.paragraphs` alone
    skips every table (found 2026-09-21): an invoice, a CV, a timetable —
    any DOCX whose content lives in cells — indexed as nearly empty."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _table_text(table: Table) -> str:
    """One line per row, cells separated by ' | ', so a row's values stay
    together in one chunk and a header row's column names are searchable."""
    lines = []
    for row in table.rows:
        cells = []
        previous = None
        for cell in row.cells:
            # Merged cells repeat the same underlying cell; keep one copy.
            if cell._tc is previous:
                continue
            previous = cell._tc
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            if text:
                cells.append(text)
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_docx(path: Path) -> list[ExtractedBlock]:
    """One block per non-empty paragraph and one per table, tagged with the
    nearest preceding heading structure: Heading 1 tracked as 'section',
    anything below that (Heading 2+) tracked as 'heading' — per the PRD's
    "paragraph, heading, section where possible" note.
    """
    doc = DocxDocument(str(path))
    blocks = []
    current_section: str | None = None
    current_heading: str | None = None

    for item in _body_items(doc):
        if isinstance(item, Table):
            text = _table_text(item)
            if text:
                blocks.append(ExtractedBlock(text=text, page_number=None, heading=current_heading, section=current_section))
            continue

        text = item.text.strip()
        if not text:
            continue

        style_name = (item.style.name if item.style else "") or ""
        if style_name == "Heading 1" or style_name == "Title":
            current_section = text
            current_heading = None
            continue
        if style_name.startswith("Heading"):
            current_heading = text
            continue

        blocks.append(
            ExtractedBlock(text=text, page_number=None, heading=current_heading, section=current_section)
        )

    return blocks
