"""PowerPoint (.pptx) extraction. Added 2026-09-20.

One block per slide, with the slide number as the block's page_number
(so results can say "Found on slide 4" the way PDFs say "page 4") and
the slide title as its heading. Speaker notes are appended to the
slide's text — they often hold the actual explanation the bullets only
hint at.
"""

from pathlib import Path

from pptx import Presentation

from .blocks import ExtractedBlock


def _shape_texts(shapes):
    for shape in shapes:
        # Grouped shapes nest; recurse so text inside a group isn't lost.
        if hasattr(shape, "shapes"):
            yield from _shape_texts(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                line = "".join(run.text for run in paragraph.runs).strip()
                if line:
                    yield line
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                line = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if line:
                    yield line


def extract_pptx(path: Path) -> list[ExtractedBlock]:
    presentation = Presentation(str(path))
    blocks = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip() or None

        lines = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)

        text = "\n".join(lines)
        if text.strip():
            blocks.append(ExtractedBlock(text=text, page_number=number, heading=title))
    return blocks
